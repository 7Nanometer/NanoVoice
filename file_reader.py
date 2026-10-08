"""
文件读取模块 —— 把用户拖进来的文章文件，读成纯文字。

支持的格式：
    .txt  / .md  / .markdown   纯文本，直接读
    .docx                      Word 文档，借助 python-docx 库提取

一个文件算一篇文章。拖进来几个文件，系统就融合几篇的内容。

为什么单独一个模块：读文件这件事看着简单，其实有两个坑——
    1. 中文 txt 的编码不统一：老文件多是 GBK，新文件多是 UTF-8，
       读错了整篇都是乱码，得挨个试。
    2. Word 文档不是纯文字，里面是 XML 结构，需要专门的库来提取。
把这两件事收在一个文件里，主程序就干净了。
"""

from pathlib import Path

import config
from logger import get_logger

logger = get_logger(__name__)

# 界面允许拖入的格式（gradio 和校验逻辑都用这个）
SUPPORTED_EXTENSIONS = [".txt", ".md", ".markdown", ".docx"]

# 试着按这些编码依次读取，哪个能读通用哪个。
# utf-8-sig 能正确处理带 BOM 的 UTF-8（Windows 记事本另存为 UTF-8 时会产生）
# gbk 是中文 Windows 的传统编码，老文件基本都是它
_ENCODINGS = ["utf-8", "utf-8-sig", "gbk", "gb18030", "big5", "latin-1"]


def read_one_file(path) -> tuple:
    """
    读取单个文件，返回 (是否成功, 文字内容或错误说明)。
    """
    p = Path(path)

    if not p.exists():
        return False, "文件不存在（可能被移动或删除了）"

    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return False, (f"不支持的格式 {suffix}。"
                       f"目前支持：{'、'.join(SUPPORTED_EXTENSIONS)}")

    try:
        if suffix == ".docx":
            text = _read_docx(p)
        else:
            text = _read_plain_text(p)
    except Exception as e:
        logger.exception("读取文件失败：%s", p)
        return False, f"读取失败：{e}"

    text = text.strip()
    if not text:
        return False, "文件里没有文字内容（可能是空文件，或扫描件图片）"

    # 太长的文件截断，避免一次喂给大模型的内容过多
    if len(text) > config.FETCH_MAX_CHARS:
        logger.info("文件 %s 较长（%s 字），截取前 %s 字",
                    p.name, len(text), config.FETCH_MAX_CHARS)
        text = text[:config.FETCH_MAX_CHARS]

    return True, text


def _read_plain_text(path: Path) -> str:
    """
    读纯文本文件，自动尝试各种编码。

    做法：挨个编码试着读一遍，哪个不报错就用哪个。
    最后用 latin-1 兜底——它永远不会失败（任何字节都能映射成字符），
    虽然中文可能是乱码，但至少不会让程序崩掉。
    """
    raw = path.read_bytes()

    # 先看开头几个字节，识别 BOM（字节顺序标记，用来标识文件编码）
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig", errors="replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        # UTF-16 的 BOM。Windows 记事本"另存为 Unicode"会产生这种文件，
        # 直接按 utf-8 读会得到一堆带空格的乱码。
        return raw.decode("utf-16", errors="replace")

    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue

    return raw.decode("utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    """
    读 Word 文档（.docx）。

    Word 文档内部是一堆 XML 和图片打包成的压缩包，不是纯文字。
    python-docx 这个库能帮我们把里面的段落和表格提取出来。

    注意：只能读 .docx（新版 Word 格式），不能读 .doc（2003 老格式）。
    """
    from docx import Document

    doc = Document(str(path))
    parts = []

    # 正文段落
    for para in doc.paragraphs:
        line = para.text.strip()
        if line:
            parts.append(line)

    # 表格里的文字也要（很多人把资料整理在表格里）
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                # 去掉相邻的重复单元格（合并单元格时会出现重复）
                deduped = [cells[0]] + [c for i, c in enumerate(cells[1:], 1)
                                        if c != cells[i - 1]]
                parts.append(" | ".join(deduped))

    return "\n".join(parts)


def read_article_files(files) -> tuple:
    """
    批量读取用户拖进来的文件。

    files 是界面上传回来的文件路径列表（gradio 传给我们的是这种形式）。
    返回 (文章文字列表, 提示信息列表)。

    例：拖进 3 个文件，其中 1 个读失败 →
        返回 (['文章1正文', '文章2正文'], ['第3个文件 xxx 读取失败：...'])
    """
    if not files:
        return [], []

    # 界面上传的可能是单个路径字符串，也可能是列表，统一成列表处理
    if isinstance(files, (str, Path)):
        files = [files]

    articles, problems = [], []
    for i, f in enumerate(files, 1):
        name = Path(f).name
        ok, result = read_one_file(f)
        if ok:
            articles.append(result)
            logger.info("已读取文件 %s（%s 字）", name, len(result))
        else:
            problems.append(f"第{i}个文件「{name}」{result}")
            logger.warning("文件读取失败 %s：%s", name, result)

    return articles, problems


def describe_files(files) -> str:
    """
    生成"已载入文章"的预览文字，显示在界面上让用户确认读到了什么。
    """
    if not files:
        return "还没有拖入文件。把 .txt / .md / .docx 文件拖到上面的框里即可，可一次拖多个。"

    if isinstance(files, (str, Path)):
        files = [files]

    lines = []
    for i, f in enumerate(files, 1):
        p = Path(f)
        ok, result = read_one_file(f)
        if ok:
            lines.append(f"✅ {i}. {p.name} —— {len(result)} 字")
        else:
            lines.append(f"❌ {i}. {p.name} —— {result}")

    good = sum(1 for line in lines if line.startswith("✅"))
    header = f"共 {len(lines)} 个文件，成功读取 {good} 个：\n"
    return header + "\n".join(lines)
